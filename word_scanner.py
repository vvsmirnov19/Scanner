import re

import docx
import openpyxl as op
import pandas as pd

from api import api_request

doc = docx.Document("tekstovye_zadachi_po_matematike.docx")
with open('text.md', 'w', encoding='utf-8') as file:
    for paragraph in doc.paragraphs:
        file.write(paragraph.text+'\n')
with open('text.md', 'r', encoding='utf-8') as file:
    text = file.readlines()
tasks = text[2:41] + text[42:97] + text[98:125] + text[126:145] + text[146:165] + text[166:190] + text[191:212] + text[213:242] + text[244:288] + text[290:310] + text[311:341] + text[342:373] + text[374:406] + text[407:449] + text[451:515] + text[516:536] + text[538:584] + text[585:608] + text[609:622] + text[623:663] + text[665:691] + text[692:760] + text[761:790] + text[791:913]
answers = text[914:940]
content = text[942:]
re_parts = re.compile(r'^(\d\.\s[А-Яа-я\s]+)\d+')
re_headers = re.compile(r'^(\d\.\d\.\s[А-Яа-я\s\«\»\,\.]+)\s*\d+')
re_tasks = re.compile(r'^(\d+\.)*\s*\t*([абв\d])*\)*\t*(.+)')
re_answer_delimeter = re.compile(r'(\d{1,4}\.)(.*?(?=\.))')
re_answer_delimeter_2 = re.compile(r'[\s\t]([\dабвгд])\)[\s\t]*(.*?(?=\;|$))')
counter = 1
header_pointer = 0
content_list = list()
for line in content:
    if re.findall(re_parts, line):
        content_list.append(
            dict(id=counter,
                 name=re.findall(re_parts, line)[0],
                 parent=0)
        )
        header_pointer = counter
    elif re.findall(re_headers, line):
        content_list.append(
            dict(
                id=counter,
                name=re.findall(re_headers, line)[0],
                parent=header_pointer)
        )
    counter += 1
df1 = pd.DataFrame.from_dict(content_list)
answ_list = list()
texter = ' '.join([answer.rstrip() for answer in answers])
tans = re.findall(re_answer_delimeter, texter)
for t in tans:
    ret = re.findall(re_answer_delimeter_2, t[1])
    if len(ret) != 0:
        for reti in ret:
            answ_list.append((f'{t[0]}{reti[0]}', reti[1]))
    else:
        answ_list.append(t)
task_pointer = None
task_list = list()
for line in tasks:
    match = re.match(re_tasks, line)
    if match:
        task_list.append(
            dict(
                id_tasks_book=f'{task_pointer if match.group(1) is None else match.group(1)}{'' if match.group(2) is None else match.group(2)}',
                task=match.group(3),
                classes='5;6',
                level=1
            )
        )
        task_pointer = task_pointer if match.group(1) is None else match.group(1)
for task in task_list:
    for answ in answ_list:
        if str(task['id_tasks_book']) == str(answ[0]) and str(task['id_tasks_book']) != '5.':
            task['answer'] = answ[1]
            break
for i in range(len(task_list)):
    if i < 39:
        task_list[i]['paragraph'] = 2
    elif i < 94:
        task_list[i]['paragraph'] = 3
    elif i < 121:
        task_list[i]['paragraph'] = 4
    elif i < 142:
        task_list[i]['paragraph'] = 5
    elif i < 159:
        task_list[i]['paragraph'] = 6
    elif i < 182:
        task_list[i]['paragraph'] = 7
    elif i < 233:
        task_list[i]['paragraph'] = 8
    elif i < 278:
        task_list[i]['paragraph'] = 10
    elif i < 297:
        task_list[i]['paragraph'] = 11
    elif i < 327:
        task_list[i]['paragraph'] = 12
    elif i < 358:
        task_list[i]['paragraph'] = 13
    elif i < 390:
        task_list[i]['paragraph'] = 14
    elif i < 432:
        task_list[i]['paragraph'] = 15
    elif i < 496:
        task_list[i]['paragraph'] = 17
    elif i < 516:
        task_list[i]['paragraph'] = 18
    elif i < 562:
        task_list[i]['paragraph'] = 20
    elif i < 585:
        task_list[i]['paragraph'] = 21
    elif i < 598:
        task_list[i]['paragraph'] = 22
    elif i < 638:
        task_list[i]['paragraph'] = 23
    elif i < 664:
        task_list[i]['paragraph'] = 25
    elif i < 732:
        task_list[i]['paragraph'] = 26
    elif i < 761:
        task_list[i]['paragraph'] = 27
    else:
        task_list[i]['paragraph'] = 28
for task in task_list[:10]:
    print(task)
    if 'answer' in task:
        task['solution'] = api_request(task['task'])
df2 = pd.DataFrame.from_dict(task_list)
with pd.ExcelWriter('table.xlsx', engine='xlsxwriter') as writer:
    df2.to_excel(writer, sheet_name='tasks', index=False)
    df1.to_excel(writer, sheet_name='table_of_contents', index=False)
wb = op.load_workbook('table.xlsx')
ws = wb.create_sheet('author')
ws.append(['name', 'author', 'description', 'topic_id', 'classes'])
ws.append(['Текстовые задачи по математике. 5–6 классы / А. В. Шевкин. — 3-е изд., перераб. — М. : Илекса, 2024. — 160 с. : ил.', 'А. В. Шевкин.', 'Сборник включает текстовые задачи по разделам школьной математики: натуральные числа, дроби, пропорции, проценты, уравнения. Ко многим задачам даны ответы или советы с чего начать решения. Решения некоторых задач приведены в качестве образцов в основном тексте книги или в разделе «Ответы, советы, решения». Материалы сборника можно использовать как дополнение к любому действующему учебнику. При подготовке этого издания добавлены новые задачи и решения некоторых задач. Пособие предназначено для учащихся 5–6 классов общеобразовательных школ, учителей, студентов педагогических вузов.', '1', '5;6'])
wb.save('table.xlsx')
