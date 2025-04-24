import re

import docx
from pdf2image import convert_from_path
from PIL import Image
import pytesseract

poppler_path = r'C:\\Program Files\\poppler\\Library\bin'

def convert_to_images(input_file, output_file, page):
    images = convert_from_path(
        input_file,
        first_page=page,
        last_page=page,
        poppler_path=poppler_path
    )
    image = images[0]
    image.save(output_file, "PNG")

def teseract_recognition(*path_img):
    text = list()
    page_gap = 40
    for image in path_img:
        img = Image.open(image)
        size = img.size[0] + page_gap
        for i in range(2):
            cropped = img.crop((i*size/2, 0, (i+1)*size/2, img.size[1]))
            cropped.save(f'cropped{i}.jpg')
            text.append(
                pytesseract.image_to_string(
                    Image.open(f'cropped{i}.jpg'),
                    lang='rus+eng',
                    config=r'--oem 3 --psm 6'
                    )
            )
        page_gap = -40
    return '\n'.join(text)

def save_text(text, name):
    with open(name, 'w', encoding='utf-8') as file:
        file.write(text)
    return

def main():
    path_name = r'C:\\Dev\\scanner\\3800.pdf'
    convert_to_images(path_name, '670.png', 670)
    convert_to_images(path_name, '671.png', 671)
    md_path = path_name[:-4] + '.md'
    save_text(teseract_recognition('670.png', '671.png'), md_path)
    re_parts = re.compile(r'\d\s([А-ЯA-Z\s]{5,})\s+')
    re_headers = re.compile(r'\d+\.([А-ЯA-Zа-яa-z\s\-]+)\.{,1}')
    re_section = re.compile(r'\d+\s([А-ЯA-Zа-яa-z\s\-]+)\.*\s+')
    result = list()
    with open(md_path, 'r', encoding='utf-8') as file:
        text = ' '.join([line.rstrip() for line in file.readlines()])
        split = re.split(re_parts, text)[1:]
        for i in range(1, len(split), 2):
            split[i] = re.split(re_headers, split[i])[1:]
            for j in range(1, len(split[i]), 2):
                split[i][j] = re.findall(re_section, split[i][j])
        result = list()
        for i in split:
            if isinstance(i, list):
                for j in i:
                    if isinstance(j, list):
                        for k in j:
                            result.append((k, '2'))
                    else:
                        result.append((j, '1'))
            else:
                result.append((i, '0'))
    result = result[:-1]
    doc = docx.Document()
    table = doc.add_table(rows=len(list(result)), cols=3)
    for row in range(len(list(result))):
        table.cell(row, 0).text = str(row+1)
        table.cell(row, 1).text = result[row][0]
        table.cell(row, 2).text = result[row][1]
    doc.save('table.docx')



if __name__ == "__main__":
    main()
